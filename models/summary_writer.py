# -*- coding: utf-8 -*-
"""
汇总Excel写入器模型
负责将处理后的数据汇总成合同开票信息提取格式
"""

import os
import logging
import shutil
from datetime import datetime
from typing import List, Dict, Any, Optional
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from collections import defaultdict


class SummaryWriter:
    """汇总Excel文件写入器"""
    
    def __init__(self):
        self.logger = logging.getLogger(__name__)
        
    def create_summary_file(self, 
                           all_data: List[Dict[str, Any]], 
                           output_dir: str, 
                           original_filename: str) -> Optional[str]:
        """
        创建汇总文件 - 合同开票信息提取格式
        
        Args:
            all_data: 所有数据列表
            output_dir: 输出目录
            original_filename: 原始文件名
            
        Returns:
            str: 输出文件路径，失败返回None
        """
        try:
            if not all_data:
                self.logger.warning("没有数据需要汇总")
                return None
                
            self.logger.info(f"开始创建汇总文件，收到 {len(all_data)} 条数据，输出目录: {output_dir}")
            
            # 调试信息：打印前几条数据的结构
            if all_data:
                sample_data = all_data[0]
                self.logger.info(f"数据样本字段: {list(sample_data.keys())}")
                self.logger.info(f"数据样本内容: {sample_data}")
            
            # 确保输出目录存在
            os.makedirs(output_dir, exist_ok=True)
            
            # 生成汇总文件名
            summary_filename = self._generate_summary_filename(original_filename)
            summary_path = os.path.join(output_dir, summary_filename)
            
            # 汇总数据
            summary_data = self._aggregate_data(all_data)
            
            if not summary_data:
                self.logger.warning("汇总后没有有效数据")
                return None
            
            self.logger.info(f"数据转换完成，生成 {len(summary_data)} 个项目")
            
            # 创建Excel文件
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "合同开票信息提取"
            
            # 写入标题
            self._write_title(worksheet)
            
            # 写入表头
            self._write_headers(worksheet)
            
            # 写入汇总数据
            self._write_summary_data(worksheet, summary_data)
            
            # 写入合计行
            self._write_total_row(worksheet, summary_data)
            
            # 设置样式和列宽
            self._format_worksheet(worksheet)
            
            # 保存文件
            workbook.save(summary_path)
            workbook.close()
            
            self.logger.info(f"成功创建汇总文件: {summary_path}")
            return summary_path
            
        except Exception as e:
            self.logger.error(f"创建汇总文件失败: {e}")
            return None
    
    def _generate_summary_filename(self, original_filename: str) -> str:
        """生成汇总文件名"""
        name_without_ext = os.path.splitext(original_filename)[0]
        timestamp = datetime.now().strftime('%Y-%m-%d')
        return f"{name_without_ext}_汇总_{timestamp}.xlsx"
    
    def _aggregate_data(self, all_data: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """
        转换数据 - 逐条转换，不合并同类项目
        """
        try:
            summary_list = []
            
            for row in all_data:
                # 品名及规格字段映射（使用国韩报税的字段）
                item_name = self._get_item_name(row)
                unit = self._get_unit(row)
                
                # 数量和金额 - 使用国韩报税的字段
                quantity = self._safe_float(row.get('quantity', 0))  # 数量字段
                amount = self._safe_float(row.get('amount', 0))     # 金额字段
                
                if item_name and (quantity > 0 or amount > 0):  # 有品名且有数量或金额
                    # 计算单价
                    unit_price = 0
                    if quantity > 0:
                        unit_price = amount / quantity
                    
                    summary_list.append({
                        'item_name': item_name,
                        'unit': unit if unit else self._get_default_unit(),
                        'quantity': int(quantity) if quantity == int(quantity) else round(quantity, 2),
                        'unit_price': round(unit_price, 2),
                        'total_amount': round(amount, 2)
                    })
            
            self.logger.info(f"数据转换完成，共 {len(summary_list)} 个项目")
            return summary_list
            
        except Exception as e:
            self.logger.error(f"数据转换失败: {e}")
            return []
    
    def _get_item_name(self, row: Dict[str, Any]) -> str:
        """获取品名及规格"""
        # 使用国韩报税的字段：product_name
        product_name = row.get('product_name', '')
        
        if product_name and str(product_name).strip():
            return str(product_name).strip()
        else:
            return "未知商品"
    
    def _get_unit(self, row: Dict[str, Any]) -> str:
        """获取单位"""
        # 使用国韩报税的字段：unit
        unit = row.get('unit', '')
        if unit and str(unit).strip():
            return str(unit).strip()
        return "个"  # 默认单位
    
    def _get_default_unit(self) -> str:
        """获取默认单位"""
        return "个"
    
    def _safe_float(self, value: Any) -> float:
        """安全转换为浮点数"""
        if value is None:
            return 0.0
        try:
            if isinstance(value, str):
                value = value.replace(',', '').replace('，', '').strip()
                if value == '' or value == '-':
                    return 0.0
            return float(value)
        except (ValueError, TypeError):
            return 0.0
    
    def _write_title(self, worksheet) -> None:
        """写入标题"""
        worksheet.merge_cells('A1:F1')
        title_cell = worksheet['A1']
        title_cell.value = "合同开票信息提取"
        title_cell.font = Font(bold=True, size=16)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[1].height = 30
    
    def _write_headers(self, worksheet) -> None:
        """写入表头"""
        headers = {
            'A2': '序号',
            'B2': '品名及规格', 
            'C2': '单位',
            'D2': '数量',
            'E2': '单价',
            'F2': '金额'
        }
        
        # 表头样式
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for cell_ref, header_text in headers.items():
            cell = worksheet[cell_ref]
            cell.value = header_text
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        worksheet.row_dimensions[2].height = 25
    
    def _write_summary_data(self, worksheet, summary_data: List[Dict[str, Any]]) -> int:
        """写入汇总数据"""
        start_row = 3
        
        # 数据样式 - 所有内容都居中显示
        data_font = Font(size=11)
        data_alignment_center = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for i, item in enumerate(summary_data):
            row_num = start_row + i
            
            # 序号
            seq_cell = worksheet.cell(row=row_num, column=1, value=i + 1)
            seq_cell.font = data_font
            seq_cell.alignment = data_alignment_center
            seq_cell.border = thin_border
            
            # 品名及规格
            name_cell = worksheet.cell(row=row_num, column=2, value=item['item_name'])
            name_cell.font = data_font
            name_cell.alignment = data_alignment_center
            name_cell.border = thin_border
            
            # 单位
            unit_cell = worksheet.cell(row=row_num, column=3, value=item['unit'])
            unit_cell.font = data_font
            unit_cell.alignment = data_alignment_center
            unit_cell.border = thin_border
            
            # 数量
            qty_cell = worksheet.cell(row=row_num, column=4, value=item['quantity'])
            qty_cell.font = data_font
            qty_cell.alignment = data_alignment_center
            qty_cell.border = thin_border
            
            # 单价
            price_cell = worksheet.cell(row=row_num, column=5, value=item['unit_price'])
            price_cell.font = data_font
            price_cell.alignment = data_alignment_center
            price_cell.border = thin_border
            
            # 金额
            amount_cell = worksheet.cell(row=row_num, column=6, value=item['total_amount'])
            amount_cell.font = data_font
            amount_cell.alignment = data_alignment_center
            amount_cell.border = thin_border
            
            worksheet.row_dimensions[row_num].height = 20
        
        return start_row + len(summary_data) - 1
    
    def _write_total_row(self, worksheet, summary_data: List[Dict[str, Any]]) -> None:
        """写入合计行"""
        if not summary_data:
            return
            
        total_row = 3 + len(summary_data)
        
        # 计算合计
        total_quantity = sum(item['quantity'] for item in summary_data)
        total_amount = sum(item['total_amount'] for item in summary_data)
        
        # 合计样式 - 所有内容都居中显示
        total_font = Font(bold=True, size=12)
        total_fill = PatternFill(start_color="FFE4B5", end_color="FFE4B5", fill_type="solid")
        total_alignment_center = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 合计标签
        total_label_cell = worksheet.cell(row=total_row, column=2, value="合计")
        total_label_cell.font = total_font
        total_label_cell.fill = total_fill
        total_label_cell.alignment = total_alignment_center
        total_label_cell.border = thin_border
        
        # 数量合计
        qty_total_cell = worksheet.cell(row=total_row, column=4, value=total_quantity)
        qty_total_cell.font = total_font
        qty_total_cell.fill = total_fill
        qty_total_cell.alignment = total_alignment_center
        qty_total_cell.border = thin_border
        
        # 金额合计
        amount_total_cell = worksheet.cell(row=total_row, column=6, value=round(total_amount, 2))
        amount_total_cell.font = total_font
        amount_total_cell.fill = total_fill
        amount_total_cell.alignment = total_alignment_center
        amount_total_cell.border = thin_border
        
        # 其他空白单元格也要设置样式
        for col in [1, 3, 5]:
            empty_cell = worksheet.cell(row=total_row, column=col, value="")
            empty_cell.font = total_font
            empty_cell.fill = total_fill
            empty_cell.alignment = total_alignment_center
            empty_cell.border = thin_border
        
        worksheet.row_dimensions[total_row].height = 25
    
    def _format_worksheet(self, worksheet) -> None:
        """设置工作表格式和列宽"""
        # 设置列宽
        column_widths = {
            'A': 8,   # 序号
            'B': 25,  # 品名及规格
            'C': 8,   # 单位
            'D': 12,  # 数量
            'E': 15,  # 单价
            'F': 15,  # 金额
        }
        
        for col, width in column_widths.items():
            worksheet.column_dimensions[col].width = width 